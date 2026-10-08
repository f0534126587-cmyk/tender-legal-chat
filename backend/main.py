from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from pydantic import BaseModel
from typing import Optional
import anthropic
import json
import uuid
import tempfile
import os
from dotenv import load_dotenv

load_dotenv()

app = FastAPI(title="Legal Tender Assistant API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

async_client = anthropic.AsyncAnthropic()

# In-memory session storage: session_id -> list of message dicts
sessions: dict[str, list[dict]] = {}

SYSTEM_PROMPT = """אתה עוזר משפטי מומחה המתמחה בעריכת מכרזים לרשויות מקומיות בישראל. תפקידך לסייע לעורכי דין בהכנת מכרזים מקצועיים, מקיפים ומדויקים מבחינה משפטית.

**חשוב ביותר: תענה תמיד בעברית בלבד. כל התגובות, השאלות והטיוטות שלך חייבות להיות בעברית.**

עקרונות פעולה:

**דיוק וסמכות**: אל תמציא תבניות. עליך להתבסס אך ורק על חוק המכרזים הישראלי (חוק הרשויות המקומיות, תקנות הרשויות המקומיות, פסיקה רלוונטית) ועל מסמכים שסיפק המשתמש. אם קיים ספק משפטי, עליך לציין זאת במפורש ולא להמציא פתרון.

**התאמה לסוג המכרז**: בתחילת התהליך, עליך לזהות או לשאול על סוג המכרז הספציפי (למשל: מכרז פומבי, מכרז זוטא, פטור ממכרז) וליישם את המסגרת המשפטית והנהלתית המחייבת החלה על אותו סוג.

**איסוף נתונים**: אם חסרים נתונים קריטיים לניסוח (כגון אופי השירות, משך החוזה, תנאי סף ספציפיים, או קריטריוני הערכה), עליך לעצור ולשאול את המשתמש בצורה ממוקדת. חל עליך איסור מוחלט לנחש או להשלים נתונים חסרים ביוזמתך.

**בהירות וסדר**: המכרז המנוסח חייב להיות בנוי בצורה היררכית, ברורה וקריאה, במטרה למזער ככל האפשר שאלות הבהרה מצד מציעים.

**מבנה המכרז** - כל מסמך חייב לכלול:
1. הזמנה להגיש הצעות
2. תנאי סף (משפטיים, טכניים וכלכליים)
3. תיאור השירותים/הטובין המבוקשים (מפרט טכני)
4. קריטריונים לבחירת הזוכה
5. טיוטת הסכם/חוזה
6. נספחים חובה (ערבויות, הצהרות מציע וכו')
7. פרטי קשר ברורים לפניות ולשאלות הבהרה

**תהליך עבודה**:
1. הערכה ראשונית: שאל את המשתמש על הגדרת המכרז הבסיסית ובקש מסמכי עזר אם נדרש.
2. ניתוח משפטי: בצע ניתוח משפטי של הדרישות.
3. אישור ניסוח: הצג טיוטה ראשונית או מתווה מפורט לאישור.
4. גרסה סופית: לאחר אישור התוכן, הצג את המכרז המלא.

**סגנון**: שפה משפטית, פורמלית, תמציתית ומקצועית ביותר בעברית."""


class ChatRequest(BaseModel):
    session_id: Optional[str] = None
    message: str


class ExportRequest(BaseModel):
    session_id: str


@app.post("/api/chat")
async def chat_stream(request: ChatRequest):
    session_id = request.session_id or str(uuid.uuid4())

    if session_id not in sessions:
        sessions[session_id] = []

    sessions[session_id].append({"role": "user", "content": request.message})

    async def generate():
        full_text = ""

        yield f"data: {json.dumps({'type': 'session_id', 'session_id': session_id}, ensure_ascii=False)}\n\n"

        try:
            async with async_client.messages.stream(
                model="claude-opus-4-8",
                max_tokens=16000,
                system=SYSTEM_PROMPT,
                messages=sessions[session_id],
            ) as stream:
                async for text in stream.text_stream:
                    full_text += text
                    yield f"data: {json.dumps({'type': 'text', 'content': text}, ensure_ascii=False)}\n\n"

            sessions[session_id].append({"role": "assistant", "content": full_text})
            yield f"data: {json.dumps({'type': 'done'})}\n\n"

        except Exception as e:
            # Revert the user message if streaming never started
            if not full_text:
                sessions[session_id].pop()
            yield f"data: {json.dumps({'type': 'error', 'message': str(e)}, ensure_ascii=False)}\n\n"

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


@app.post("/api/export")
async def export_to_word(request: ExportRequest):
    if request.session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")

    messages = sessions[request.session_id]
    if not messages:
        raise HTTPException(status_code=400, detail="No messages to export")

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        def make_rtl(paragraph):
            pPr = paragraph._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            pPr.insert(0, bidi)

        doc = Document()

        # Title
        title = doc.add_heading("עוזר משפטי - מכרזים", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        make_rtl(title)

        doc.add_paragraph()

        for msg in messages:
            role_label = "עורך דין" if msg["role"] == "user" else "עוזר משפטי"
            content = msg["content"]

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            make_rtl(p)
            run = p.add_run(f"{role_label}: ")
            run.bold = True

            # Split content by newlines for proper paragraph formatting
            lines = content.split("\n")
            p.add_run(lines[0])

            for line in lines[1:]:
                extra = doc.add_paragraph(line)
                extra.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                make_rtl(extra)

            doc.add_paragraph()

        tmp_file = tempfile.mktemp(suffix=".docx")
        doc.save(tmp_file)

        return FileResponse(
            tmp_file,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="tender_draft.docx",
        )

    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx not installed. Run: pip install python-docx")


@app.get("/api/session/{session_id}")
async def get_session(session_id: str):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="Session not found")
    return {"session_id": session_id, "messages": sessions[session_id]}


@app.delete("/api/session/{session_id}")
async def clear_session(session_id: str):
    if session_id in sessions:
        del sessions[session_id]
    return {"status": "cleared"}


@app.get("/api/health")
async def health_check():
    return {"status": "ok", "model": "claude-opus-4-8"}
